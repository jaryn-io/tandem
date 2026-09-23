"""Tests for the S01 core import engine. Run with: python3 -m unittest discover -s tests"""

import json
import sqlite3
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from pricelist_import import (
    AmbiguousPriceError,
    CatalogStore,
    NumericCell,
    ProfileError,
    ReaderError,
    apply_import,
    load_profile,
    normalize_code,
    parse_file,
    parse_price,
    validate_row,
)
from pricelist_import.profiles import profile_from_dict
from unittest import mock


def make_profile(**overrides) -> dict:
    base = {
        "supplier_id": "acme",
        "format": "csv",
        "csv": {"delimiter": ";", "encoding": "utf-8"},
        "columns": {
            "article_code": "Codice",
            "description": "Descrizione",
            "price": "Prezzo",
        },
        "decimal": {"decimal_separator": ",", "thousands_separator": "."},
        "code_normalization": {"strip_prefix": "AC-", "uppercase": True},
        "code_pattern": "^[A-Z0-9]+$",
        "defaults": {"currency": "EUR"},
    }
    base.update(overrides)
    return base


class TestProfiles(unittest.TestCase):
    def test_loads_minimal_profile(self):
        profile = load_profile(make_profile())
        self.assertEqual(profile.supplier_id, "acme")
        self.assertEqual(profile.csv.delimiter, ";")
        self.assertEqual(profile.decimal.decimal_separator, ",")
        self.assertEqual(profile.defaults["currency"], "EUR")

    def test_loads_from_json_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "acme.json"
            path.write_text(json.dumps(make_profile()), encoding="utf-8")
            profile = load_profile(path)
            self.assertEqual(profile.format, "csv")

    def test_rejects_bad_supplier_id(self):
        with self.assertRaises(ProfileError):
            profile_from_dict(make_profile(supplier_id="Bad ID!"))

    def test_rejects_unknown_format(self):
        with self.assertRaises(ProfileError):
            profile_from_dict(make_profile(format="xml"))

    def test_rejects_missing_required_column(self):
        profile = make_profile()
        del profile["columns"]["price"]
        with self.assertRaises(ProfileError):
            profile_from_dict(profile)

    def test_rejects_unknown_column_key(self):
        profile = make_profile()
        profile["columns"]["weight"] = "Peso"
        with self.assertRaises(ProfileError):
            profile_from_dict(profile)

    def test_rejects_equal_separators(self):
        profile = make_profile()
        profile["decimal"] = {"decimal_separator": ",", "thousands_separator": ","}
        with self.assertRaises(ProfileError):
            profile_from_dict(profile)

    def test_rejects_bad_regex(self):
        with self.assertRaises(ProfileError):
            profile_from_dict(make_profile(code_pattern="^(unclosed"))


class TestValidation(unittest.TestCase):
    def setUp(self):
        self.profile = load_profile(make_profile())

    def test_normalize_code_strips_prefix_and_uppercases(self):
        self.assertEqual(
            normalize_code(" ac-ab 12 ", self.profile.code_normalization), "AB12"
        )

    def test_parse_italian_decimal(self):
        self.assertEqual(
            parse_price("1.234,56", self.profile.decimal), Decimal("1234.56")
        )

    def test_parse_thousands_only_text_resolves_by_profile(self):
        # "1.234" under a "."-thousands profile means one thousand two hundred
        # thirty-four, not 1.234.
        self.assertEqual(parse_price("1.234", self.profile.decimal), Decimal("1234"))
        self.assertEqual(
            parse_price("1.234.567", self.profile.decimal), Decimal("1234567")
        )

    def test_parse_numeric_cell_passthrough(self):
        # A numeric Excel cell arrives machine-formatted and marked; it is
        # accepted verbatim whatever the profile's separators are.
        self.assertEqual(
            parse_price(NumericCell("42.5"), self.profile.decimal), Decimal("42.5")
        )

    def test_parse_ambiguous_text_price_is_rejected(self):
        # Textual "42.5" is ambiguous here: "." is the thousands separator and
        # there is no decimal comma. It must not silently become 42.5 or 425.
        with self.assertRaises(AmbiguousPriceError):
            parse_price("42.5", self.profile.decimal)
        with self.assertRaises(AmbiguousPriceError):
            parse_price("12.90", self.profile.decimal)

    def test_parse_malformed_thousands_grouping_is_rejected(self):
        # Malformed grouping must not be silently repaired by removing every
        # thousands separator: these would otherwise become 123.45, 1234.56
        # and 1234.56.
        for bad in ("1.2.3,45", "12.34,56", "1..234,56"):
            with self.assertRaises(AmbiguousPriceError):
                parse_price(bad, self.profile.decimal)
        # Well-formed groups and unseparated integers still parse.
        self.assertEqual(
            parse_price("1.234,45", self.profile.decimal), Decimal("1234.45")
        )
        self.assertEqual(
            parse_price("1.234.567,89", self.profile.decimal),
            Decimal("1234567.89"),
        )
        self.assertEqual(
            parse_price("1234,45", self.profile.decimal), Decimal("1234.45")
        )
        self.assertEqual(parse_price("12,50", self.profile.decimal), Decimal("12.50"))

    def test_parse_thousands_separator_in_fraction_is_rejected(self):
        # A thousands separator inside the fractional part must be caught
        # before any separator is removed: these would otherwise silently
        # become 12.345 and 1234.56.
        for bad in ("12,3.45", "1.234,5.6"):
            with self.assertRaises(AmbiguousPriceError):
                parse_price(bad, self.profile.decimal)
        # A legitimately grouped price with a plain fraction still parses.
        self.assertEqual(
            parse_price("1.234,56", self.profile.decimal), Decimal("1234.56")
        )

    def test_validate_row_reports_fraction_separator_as_ambiguous(self):
        row, rejection = validate_row(
            5,
            {"article_code": "AC-8", "description": None, "price": "12,3.45",
             "currency": None},
            self.profile,
        )
        self.assertIsNone(row)
        self.assertEqual(rejection.reason_code, "ambiguous_price")
        self.assertIn("ambiguous price", rejection.message)

    def test_validate_row_reports_malformed_grouping_as_ambiguous(self):
        row, rejection = validate_row(
            4,
            {"article_code": "AC-7", "description": None, "price": "1.2.3,45",
             "currency": None},
            self.profile,
        )
        self.assertIsNone(row)
        self.assertEqual(rejection.reason_code, "ambiguous_price")
        self.assertIn("ambiguous price", rejection.message)

    def test_validate_row_reports_ambiguous_price_reason(self):
        row, rejection = validate_row(
            6,
            {"article_code": "AC-9", "description": None, "price": "12.90",
             "currency": None},
            self.profile,
        )
        self.assertIsNone(row)
        self.assertEqual(rejection.reason_code, "ambiguous_price")
        self.assertIn("ambiguous price", rejection.message)

    def test_parse_non_numeric_text_with_dots_is_invalid_not_ambiguous(self):
        with self.assertRaises(ValueError) as ctx:
            parse_price("n.d.", self.profile.decimal)
        self.assertNotIsInstance(ctx.exception, AmbiguousPriceError)

    def test_rejects_non_numeric_price(self):
        with self.assertRaises(ValueError):
            parse_price("n/a", self.profile.decimal)

    def test_validate_row_ok(self):
        row, rejection = validate_row(
            2,
            {"article_code": "AC-123", "description": " Vite ", "price": "12,50",
             "currency": None},
            self.profile,
        )
        self.assertIsNone(rejection)
        self.assertEqual(row.article_code, "123")
        self.assertEqual(row.description, "Vite")
        self.assertEqual(row.price, Decimal("12.50"))
        self.assertEqual(row.currency, "EUR")

    def test_validate_row_collects_multiple_errors(self):
        row, rejection = validate_row(
            3,
            {"article_code": "", "description": None, "price": "abc",
             "currency": "euro"},
            self.profile,
        )
        self.assertIsNone(row)
        self.assertEqual(rejection.reason_code, "missing_article_code")
        self.assertIn("invalid price", rejection.message)
        self.assertIn("currency", rejection.message)

    def test_validate_row_rejects_negative_price(self):
        row, rejection = validate_row(
            4,
            {"article_code": "X1", "description": None, "price": "-5,00",
             "currency": None},
            self.profile,
        )
        self.assertIsNone(row)
        self.assertEqual(rejection.reason_code, "negative_price")

    def test_validate_row_rejects_code_pattern_mismatch(self):
        row, rejection = validate_row(
            5,
            {"article_code": "AC-a?b", "description": None, "price": "1,00",
             "currency": None},
            self.profile,
        )
        self.assertIsNone(row)
        self.assertEqual(rejection.reason_code, "invalid_article_code")


class TestCsvReader(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def write_csv(self, text: str, name: str = "listino.csv") -> Path:
        path = self.dir / name
        path.write_text(text, encoding="utf-8")
        return path

    def test_reads_rows_with_physical_line_numbers(self):
        path = self.write_csv(
            "Codice;Descrizione;Prezzo\nAC-1;Vite;1,00\n\nAC-2;Bullone;2,50\n"
        )
        profile = load_profile(make_profile())
        rows = list(__import__("pricelist_import").iter_source_rows(path, profile))
        self.assertEqual([r[0] for r in rows], [2, 4])
        self.assertEqual(rows[0][1]["article_code"], "AC-1")
        self.assertIsNone(rows[0][1]["currency"])

    def test_missing_required_column_raises(self):
        path = self.write_csv("Codice;Descrizione\nAC-1;Vite\n")
        profile = load_profile(make_profile())
        with self.assertRaises(ReaderError):
            list(__import__("pricelist_import").iter_source_rows(path, profile))

    def test_wrong_encoding_raises_reader_error(self):
        path = self.dir / "latin1.csv"
        path.write_bytes("Codice;Descrizione;Prezzo\nAC-1;àèì;1,00\n".encode("latin-1"))
        profile = load_profile(make_profile())
        with self.assertRaises(ReaderError):
            list(__import__("pricelist_import").iter_source_rows(path, profile))

    def test_skip_rows(self):
        path = self.write_csv(
            "Listino prezzi 2026\nCodice;Descrizione;Prezzo\nAC-1;Vite;1,00\n"
        )
        profile_dict = make_profile()
        profile_dict["csv"]["skip_rows"] = 1
        profile = load_profile(profile_dict)
        rows = list(__import__("pricelist_import").iter_source_rows(path, profile))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], 3)


class TestEngine(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def write_csv(self, text: str) -> Path:
        path = self.dir / "listino.csv"
        path.write_text(text, encoding="utf-8")
        return path

    def test_parse_file_splits_accepted_rejected_duplicates(self):
        path = self.write_csv(
            "Codice;Descrizione;Prezzo\n"
            "AC-1;Vite;1,00\n"
            "AC-2;Bullone;non-esiste\n"
            "AC-1;Vite duplicata;2,00\n"
            ";Senza codice;3,00\n"
            "AC-3;Rondella;0,75\n"
        )
        report = parse_file(path, load_profile(make_profile()))
        self.assertEqual(report.rows_total, 5)
        self.assertEqual([r.article_code for r in report.accepted], ["1", "3"])
        reasons = {r.row_number: r.reason_code for r in report.rejections}
        self.assertEqual(reasons[3], "invalid_price")
        self.assertEqual(reasons[4], "duplicate_in_file")
        self.assertEqual(reasons[5], "missing_article_code")

    def test_catalog_duplicate_classification_is_read_only(self):
        path = self.write_csv("Codice;Descrizione;Prezzo\nAC-1;Vite;1,00\nAC-9;Nuovo;9,99\n")
        store = CatalogStore(self.dir / "catalog.db")
        run = store.begin_run("seed", "seed.csv", "acme")
        store.upsert_article(run, "1", "Vite", "0.90", "EUR", changed=True, supplier_id="acme")
        store.commit()
        report = parse_file(path, load_profile(make_profile()), store=store)
        self.assertEqual(report.catalog_duplicates, ["1"])
        self.assertEqual(len(report.accepted), 2)
        self.assertEqual(len(store.list_articles()), 1)
        store.close()

    def test_missing_file_raises_reader_error(self):
        with self.assertRaises(ReaderError):
            parse_file(self.dir / "nope.csv", load_profile(make_profile()))

    def test_file_hash_is_stable(self):
        path = self.write_csv("Codice;Descrizione;Prezzo\nAC-1;Vite;1,00\n")
        profile = load_profile(make_profile())
        first = parse_file(path, profile)
        second = parse_file(path, profile)
        self.assertEqual(first.file_sha256, second.file_sha256)
        self.assertEqual(len(first.file_sha256), 64)


class TestCrossSupplierGuard(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.store = CatalogStore(self.dir / "catalog.db")
        run = self.store.begin_run("seed", "seed.csv", "acme")
        self.store.upsert_article(
            run, "SHARED", "Original", "10.00", "EUR", changed=True,
            supplier_id="acme",
        )
        self.store.commit()

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def test_cross_supplier_code_is_rejected_and_never_applied(self):
        path = self.dir / "hostile.csv"
        path.write_text(
            "Codice;Descrizione;Prezzo\nAC-SHARED;Hostile;1,00\nAC-NEW;Legit;2,00\n",
            encoding="utf-8",
        )
        profile = load_profile(make_profile(supplier_id="other"))
        report = parse_file(path, profile, store=self.store)
        self.assertEqual([r.article_code for r in report.accepted], ["NEW"])
        self.assertEqual(
            [r.reason_code for r in report.rejections], ["cross_supplier_code"]
        )
        self.assertIn("acme", report.rejections[0].message)

        result = apply_import(report, self.store)
        self.assertEqual(result.inserted, 1)
        article = self.store.get_article("SHARED")
        self.assertEqual(article.supplier_id, "acme")
        self.assertEqual(article.current_price, "10.00")
        self.assertEqual(article.description, "Original")
        self.assertEqual(len(self.store.price_history("SHARED")), 1)


class TestInputBounds(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def write_csv(self, text: str) -> Path:
        path = self.dir / "listino.csv"
        path.write_text(text, encoding="utf-8")
        return path

    def test_oversized_source_file_is_refused(self):
        path = self.write_csv("Codice;Descrizione;Prezzo\nAC-1;Vite;1,00\n")
        with mock.patch("pricelist_import.readers.MAX_SOURCE_BYTES", 10):
            with self.assertRaises(ReaderError) as ctx:
                parse_file(path, load_profile(make_profile()))
        self.assertIn("limit", str(ctx.exception))

    def test_row_limit_is_refused_with_file_level_error(self):
        rows = "".join(f"AC-{i};Vite;1,00\n" for i in range(6))
        path = self.write_csv("Codice;Descrizione;Prezzo\n" + rows)
        with mock.patch("pricelist_import.engine.MAX_ROWS", 5):
            with self.assertRaises(ReaderError) as ctx:
                parse_file(path, load_profile(make_profile()))
        self.assertIn("data rows", str(ctx.exception))

    def test_xlsx_over_expansion_limit_is_refused(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed")
        path = self.dir / "listino.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Prezzi"
        ws.append(["Codice", "Descrizione", "Prezzo"])
        ws.append(["AC-1", "Vite", 1.5])
        wb.save(path)
        profile = load_profile(make_profile(
            format="xlsx", xlsx={"sheet": "Prezzi", "header_row": 1}
        ))
        with mock.patch(
            "pricelist_import.readers.MAX_XLSX_UNCOMPRESSED_BYTES", 100
        ):
            with self.assertRaises(ReaderError) as ctx:
                parse_file(path, profile)
        self.assertIn("decompression limit", str(ctx.exception))

    def test_xlsx_size_check_runs_before_decompressing_integrity_check(self):
        # testzip() decompresses archive members; the declared-size bound must
        # refuse an oversized archive before any decompression work happens.
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed")
        path = self.dir / "listino.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Prezzi"
        ws.append(["Codice", "Descrizione", "Prezzo"])
        ws.append(["AC-1", "Vite", 1.5])
        wb.save(path)
        profile = load_profile(make_profile(
            format="xlsx", xlsx={"sheet": "Prezzi", "header_row": 1}
        ))
        with mock.patch(
            "zipfile.ZipFile.testzip",
            side_effect=AssertionError("testzip ran before the size check"),
        ):
            with mock.patch(
                "pricelist_import.readers.MAX_XLSX_UNCOMPRESSED_BYTES", 100
            ):
                with self.assertRaises(ReaderError) as ctx:
                    parse_file(path, profile)
        self.assertIn("decompression limit", str(ctx.exception))


class TestXlsxNumericCellPrecision(unittest.TestCase):
    """Excel numeric cells keep their available precision when converted to
    text: no fixed-decimal truncation of supplier-provided prices."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def _write_xlsx(self, name: str, prices: list[float]) -> Path:
        import openpyxl

        path = self.dir / name
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Prezzi"
        ws.append(["Codice", "Descrizione", "Prezzo"])
        for i, price in enumerate(prices, start=1):
            ws.append([f"AC-{i}", "Vite", price])
        wb.save(path)
        return path

    def _profile(self):
        return load_profile(
            make_profile(format="xlsx", xlsx={"sheet": "Prezzi", "header_row": 1})
        )

    def test_high_precision_and_tiny_numeric_prices_are_preserved(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl not installed")
        path = self._write_xlsx("listino.xlsx", [1.23456789123, 0.00000000001])
        report = parse_file(path, self._profile())
        self.assertEqual(report.rows_rejected, 0)
        prices = {row.article_code: row.price for row in report.accepted}
        self.assertEqual(prices["1"], Decimal("1.23456789123"))
        self.assertEqual(prices["2"], Decimal("0.00000000001"))

    def test_precision_survives_apply_and_price_history(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl not installed")
        store = CatalogStore(self.dir / "catalog.db")
        first = parse_file(
            self._write_xlsx("v1.xlsx", [1.23456789123]), self._profile()
        )
        apply_import(first, store)
        article = store.get_article("1")
        self.assertEqual(article.current_price, "1.23456789123")

        second = parse_file(
            self._write_xlsx("v2.xlsx", [1.23456789124]), self._profile()
        )
        result = apply_import(second, store)
        self.assertEqual(result.updated, 1)
        history = store.price_history("1")
        self.assertEqual(
            [price for price, _, _, _ in history],
            ["1.23456789123", "1.23456789124"],
        )
        store.close()


class TestStore(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = CatalogStore(Path(self.tmp.name) / "catalog.db")

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def test_upsert_insert_then_change_then_unchanged(self):
        run1 = self.store.begin_run("h1", "f1.csv", "acme")
        self.store.upsert_article(run1, "A1", "Vite", "1.00", "EUR", changed=True, supplier_id="acme")
        self.store.commit()
        article = self.store.get_article("A1")
        self.assertEqual(article.current_price, "1.00")
        self.assertEqual(article.supplier_id, "acme")
        self.assertEqual(article.first_seen_run, run1)
        self.assertEqual(len(self.store.price_history("A1")), 1)

        run2 = self.store.begin_run("h2", "f2.csv", "acme")
        self.store.upsert_article(run2, "A1", "Vite", "1.20", "EUR", changed=True, supplier_id="acme")
        self.store.commit()
        self.assertEqual(self.store.get_article("A1").current_price, "1.20")
        self.assertEqual(self.store.get_article("A1").last_changed_run, run2)
        self.assertEqual(len(self.store.price_history("A1")), 2)

        run3 = self.store.begin_run("h3", "f3.csv", "acme")
        self.store.upsert_article(run3, "A1", "Vite", "1.20", "EUR", changed=False, supplier_id="acme")
        self.store.commit()
        self.assertEqual(self.store.get_article("A1").last_changed_run, run2)
        self.assertEqual(len(self.store.price_history("A1")), 2)

    def test_description_refresh_without_price_change_appends_no_history(self):
        run1 = self.store.begin_run("h1", "f1.csv", "acme")
        self.store.upsert_article(run1, "A1", "Vite", "1.00", "EUR", changed=True, supplier_id="acme")
        self.store.commit()
        run2 = self.store.begin_run("h2", "f2.csv", "acme")
        self.store.upsert_article(run2, "A1", "Vite M4", "1.00", "EUR", changed=False, supplier_id="acme")
        self.store.commit()
        article = self.store.get_article("A1")
        self.assertEqual(article.description, "Vite M4")
        self.assertEqual(article.last_changed_run, run1)
        self.assertEqual(len(self.store.price_history("A1")), 1)

    def test_list_articles_filters_by_supplier(self):
        run = self.store.begin_run("h1", "f1.csv", "acme")
        self.store.upsert_article(run, "A1", "Vite", "1.00", "EUR", changed=True, supplier_id="acme")
        self.store.upsert_article(run, "B1", "Screw", "2.00", "EUR", changed=True, supplier_id="nordwind")
        self.store.commit()
        self.assertEqual([a.article_code for a in self.store.list_articles()], ["A1", "B1"])
        self.assertEqual(
            [a.article_code for a in self.store.list_articles(supplier_id="nordwind")],
            ["B1"],
        )

    def test_find_run_by_hash_supports_idempotency_check(self):
        run = self.store.begin_run("abc123", "f.csv", "acme")
        self.store.finish_run(run, 10, 9, 1)
        self.store.commit()
        found = self.store.find_run_by_hash("abc123")
        self.assertIsNotNone(found)
        self.assertEqual(found.rows_accepted, 9)
        self.assertIsNone(self.store.find_run_by_hash("other"))

    def test_rollback_discards_uncommitted_run(self):
        run = self.store.begin_run("h", "f.csv", "acme")
        self.store.upsert_article(run, "A1", "Vite", "1.00", "EUR", changed=True, supplier_id="acme")
        self.store.rollback()
        self.assertFalse(self.store.article_exists("A1"))


class TestExcelReader(unittest.TestCase):
    def setUp(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl not installed")

    def test_reads_xlsx_by_profile(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "listino.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Prezzi"
            ws.append(["Listino 2026"])
            ws.append(["Codice", "Descrizione", "Prezzo"])
            ws.append(["AC-1", "Vite", 1.5])
            ws.append([None, None, None])
            ws.append(["AC-2", "Bullone", "2,75"])
            wb.save(path)

            profile = load_profile(make_profile(
                format="xlsx",
                xlsx={"sheet": "Prezzi", "header_row": 2},
            ))
            report = parse_file(path, profile)
            self.assertEqual(report.rows_total, 2)
            self.assertEqual([r.article_code for r in report.accepted], ["1", "2"])
            self.assertEqual(report.accepted[0].price, Decimal("1.5"))
            self.assertEqual(report.accepted[1].price, Decimal("2.75"))

    def test_missing_sheet_raises_reader_error(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "listino.xlsx"
            wb = openpyxl.Workbook()
            wb.active.append(["Codice", "Prezzo"])
            wb.save(path)
            profile = load_profile(make_profile(
                format="xlsx", xlsx={"sheet": "Inesistente"}
            ))
            with self.assertRaises(ReaderError):
                parse_file(path, profile)


if __name__ == "__main__":
    unittest.main()
